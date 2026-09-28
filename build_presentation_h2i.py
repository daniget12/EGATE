import os
import sys
import uuid
import requests
from html2image import Html2Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_test_user():
    print("Creating test user via app context...")
    sys.path.append(os.path.dirname(os.path.abspath(__file__)))
    from dotenv import load_dotenv
    load_dotenv()
    from app import create_app, db
    from app.models import User
    
    app = create_app()
    with app.app_context():
        username = f"ppt_user_{uuid.uuid4().hex[:6]}"
        password = "password"
        user = User(username=username, email=f"{username}@example.com")
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        return username, password

def take_screenshots(username, password):
    print("Taking screenshots using html2image...")
    shots = {}
    
    hti = Html2Image(browser='edge')
    hti.size = (1280, 720)
    
    # 1. Home
    print("Capturing Home...")
    hti.screenshot(url='http://127.0.0.1:5000/', save_as='shot_home.png')
    shots['home'] = 'shot_home.png'
    
    # Login via requests
    session = requests.Session()
    # get CSRF if any? No, typical flask-login doesn't enforce CSRF without flask-wtf.
    # The form just has username/password.
    login_resp = session.post("http://127.0.0.1:5000/auth/login", data={"username": username, "password": password}, allow_redirects=True)
    
    def capture_auth_page(url_path, save_name):
        print(f"Capturing {url_path}...")
        html = session.get(f"http://127.0.0.1:5000{url_path}").text
        # Fix relative paths for static assets so html2image can load them
        html = html.replace('href="/static', 'href="http://127.0.0.1:5000/static')
        html = html.replace('src="/static', 'src="http://127.0.0.1:5000/static')
        hti.screenshot(html_str=html, save_as=save_name)
        return save_name

    shots['challenges'] = capture_auth_page('/challenges/', 'shot_challenges.png')
    shots['courses'] = capture_auth_page('/courses/', 'shot_courses.png')
    shots['leaderboard'] = capture_auth_page('/leaderboard', 'shot_leaderboard.png')
    
    return shots

def build_ppt(shots):
    print("Building presentation...")
    prs = Presentation()
    TITLE_SLIDE = 0
    BULLET_SLIDE = 1
    BLANK_SLIDE = 6
    
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE])
    slide.shapes.title.text = "EGATE Talent Center"
    slide.placeholders[1].text = "Cybersecurity Learning Platform\nProject Presentation"
    
    slide = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE])
    slide.shapes.title.text = "The Objective"
    tf = slide.placeholders[1].text_frame
    tf.text = "The Need for Practical Cybersecurity Education"
    p = tf.add_paragraph()
    p.text = "Theoretical knowledge is not enough to stop modern cyber threats."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "EGATE Talent Center provides a hands-on, gamified environment."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Users learn by doing: solving real-world Capture The Flag (CTF) challenges."
    p.level = 1
    
    slide = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE])
    slide.shapes.title.text = "Platform Features"
    tf = slide.placeholders[1].text_frame
    tf.text = "Intelligent AI Helper for personalized student assistance."
    p = tf.add_paragraph()
    p.text = "Interactive CTF Challenges across multiple domains (Web, Crypto, Binary, etc.)"
    p = tf.add_paragraph()
    p.text = "Live Global Leaderboard to drive competition."
    p = tf.add_paragraph()
    p.text = "Structured Learning Paths across multiple departments."
    p = tf.add_paragraph()
    p.text = "Instant flag verification & scoring system."
    
    slide = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE])
    slide.shapes.title.text = "Upcoming Departments & Courses"
    tf = slide.placeholders[1].text_frame
    tf.text = "We are expanding beyond Cybersecurity! Releasing soon:"
    p = tf.add_paragraph()
    p.text = "Digital & Emerging Technologies"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Robotics & Electronics"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Astronomy & Space Technology"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Biotechnology"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "STEM (Science, Technology, Engineering & Mathematics)"
    p.level = 1
    
    def add_image_slide(title_text, img_path):
        slide = prs.slides.add_slide(prs.slide_layouts[BLANK_SLIDE])
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 51, 102)
        if img_path and os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(1), Inches(1.2), width=Inches(8))
    
    add_image_slide("Landing Page", shots.get('home'))
    add_image_slide("Courses & Departments", shots.get('courses'))
    add_image_slide("Challenges Dashboard", shots.get('challenges'))
    add_image_slide("Live Leaderboard", shots.get('leaderboard'))
    
    slide = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE])
    slide.shapes.title.text = "Technology Stack"
    tf = slide.placeholders[1].text_frame
    tf.text = "Backend: Python, Flask, SQLAlchemy"
    p = tf.add_paragraph()
    p.text = "Frontend: HTML5, CSS3, Bootstrap 5, Javascript"
    p = tf.add_paragraph()
    p.text = "Database: PostgreSQL (Neon Serverless)"
    p = tf.add_paragraph()
    p.text = "Deployment: Ready for Vercel"
    
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE])
    slide.shapes.title.text = "Thank You"
    slide.placeholders[1].text = "Questions & Answers"
    
    out_file = "EGATE_Project_Presentation_V2.pptx"
    prs.save(out_file)
    print(f"Presentation saved successfully to {out_file}")

if __name__ == "__main__":
    try:
        u, p = create_test_user()
        shots = take_screenshots(u, p)
        build_ppt(shots)
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"Error occurred: {e}")
