import os
import sys
import uuid
import time
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_test_user():
    print("Creating test user via app context...")
    sys.path.append(os.path.dirname(os.path.abspath(__file__)))
    from dotenv import load_dotenv
    load_dotenv()
    from app import create_app, db
    from app.models import User
    
    app = create_app()
    with app.app_context():
        username = f"ppt_user_{uuid.uuid4().hex[:6]}"
        password = "password"
        user = User(username=username, email=f"{username}@example.com")
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        return username, password

def take_screenshots(username, password):
    print("Taking screenshots...")
    shots = {}
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, channel="msedge")
        context = browser.new_context(viewport={"width": 1280, "height": 720})
        page = context.new_page()
        
        # Home
        page.goto("http://127.0.0.1:5000/")
        page.wait_for_timeout(1000)
        page.screenshot(path="shot_home.png")
        shots['home'] = "shot_home.png"
        
        # Login
        page.goto("http://127.0.0.1:5000/auth/login")
        page.fill(".sliding-sign-in-container input[name=username]", username)
        page.fill(".sliding-sign-in-container input[name=password]", password)
        page.click(".sliding-sign-in-container button[type=submit]")
        page.wait_for_url("**/challenges*")
        
        # Challenges
        page.wait_for_timeout(500)
        page.screenshot(path="shot_challenges.png")
        shots['challenges'] = "shot_challenges.png"
        
        # Challenge Detail
        try:
            # Wait for either 'View Challenge' or similar link
            page.click("a.btn:has-text('View Challenge')", timeout=3000)
            page.wait_for_timeout(500)
            page.screenshot(path="shot_challenge_detail.png")
            shots['challenge_detail'] = "shot_challenge_detail.png"
        except Exception as e:
            print("Could not click challenge detail:", e)
            shots['challenge_detail'] = "shot_challenges.png"
        
        # Leaderboard
        page.goto("http://127.0.0.1:5000/leaderboard")
        page.wait_for_timeout(500)
        page.screenshot(path="shot_leaderboard.png")
        shots['leaderboard'] = "shot_leaderboard.png"
        
        browser.close()
    return shots

def build_ppt(shots):
    print("Building presentation...")
    prs = Presentation()
    TITLE_SLIDE = 0
    BULLET_SLIDE = 1
    BLANK_SLIDE = 6
    
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE])
    slide.shapes.title.text = "EGATE Talent Center"
    slide.placeholders[1].text = "Cybersecurity Learning Platform\nProject Presentation"
    
    slide = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE])
    slide.shapes.title.text = "The Objective"
    tf = slide.placeholders[1].text_frame
    tf.text = "The Need for Practical Cybersecurity Education"
    p = tf.add_paragraph()
    p.text = "Theoretical knowledge is not enough to stop modern cyber threats."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "EGATE Talent Center provides a hands-on, gamified environment."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Users learn by doing: solving real-world Capture The Flag (CTF) challenges."
    p.level = 1
    
    slide = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE])
    slide.shapes.title.text = "Platform Features"
    tf = slide.placeholders[1].text_frame
    tf.text = "Interactive CTF Challenges across multiple domains (Web, Crypto, Binary, etc.)"
    p = tf.add_paragraph()
    p.text = "Live Global Leaderboard to drive competition."
    p = tf.add_paragraph()
    p.text = "Structured Learning Paths to guide beginners."
    p = tf.add_paragraph()
    p.text = "Instant flag verification & scoring system."
    
    def add_image_slide(title_text, img_path):
        slide = prs.slides.add_slide(prs.slide_layouts[BLANK_SLIDE])
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 51, 102)
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(1), Inches(1.2), width=Inches(8))
    
    add_image_slide("Landing Page", shots.get('home'))
    add_image_slide("Challenges Dashboard", shots.get('challenges'))
    add_image_slide("Challenge Interface", shots.get('challenge_detail'))
    add_image_slide("Live Leaderboard", shots.get('leaderboard'))
    
    slide = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE])
    slide.shapes.title.text = "Technology Stack"
    tf = slide.placeholders[1].text_frame
    tf.text = "Backend: Python, Flask, SQLAlchemy"
    p = tf.add_paragraph()
    p.text = "Frontend: HTML5, CSS3, Bootstrap 5, Javascript"
    p = tf.add_paragraph()
    p.text = "Database: PostgreSQL (Neon Serverless)"
    p = tf.add_paragraph()
    p.text = "Deployment: Ready for Vercel"
    
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE])
    slide.shapes.title.text = "Thank You"
    slide.placeholders[1].text = "Questions & Answers"
    
    out_file = "EGATE_Project_Presentation.pptx"
    prs.save(out_file)
    print(f"Presentation saved successfully to {out_file}")

if __name__ == "__main__":
    try:
        u, p = create_test_user()
        shots = take_screenshots(u, p)
        build_ppt(shots)
    except Exception as e:
        print(f"Error occurred: {e}")
