from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Slide layouts
TITLE_SLIDE_LAYOUT = 0
BULLET_SLIDE_LAYOUT = 1
SECTION_HEADER_LAYOUT = 2

# 1. Title Slide
slide_layout = prs.slide_layouts[TITLE_SLIDE_LAYOUT]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "EGATE Talent Center"
subtitle.text = "Technology Stack Overview\nAn Analysis of the Platform Architecture"

def add_slide(title_text, bullet_points):
    slide_layout = prs.slide_layouts[BULLET_SLIDE_LAYOUT]
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = title_text
    
    tf = body_shape.text_frame
    tf.text = bullet_points[0]
    
    for point in bullet_points[1:]:
        p = tf.add_paragraph()
        if point.startswith("  -"):
            p.text = point.replace("  - ", "")
            p.level = 1
        else:
            p.text = point
            p.level = 0
            
    return slide

# 2. Backend Stack
add_slide("Backend Framework", [
    "Python 3.x",
    "  - The core programming language powering the logic.",
    "Flask (Web Framework)",
    "  - A lightweight WSGI web application framework.",
    "  - Handles HTTP routing, requests, and responses.",
    "Application Factory Pattern",
    "  - Modular architecture (app/__init__.py) for scalability and testing.",
    "Blueprints",
    "  - Routes are separated into modules: auth, main, challenges, courses, etc."
])

# 3. Database & ORM
add_slide("Database & ORM", [
    "Flask-SQLAlchemy (ORM)",
    "  - Maps Python classes to database tables (User, Challenge, UserChallenge).",
    "  - Prevents SQL Injection through parameterization.",
    "Neon PostgreSQL",
    "  - The primary production database used for the platform.",
    "  - Serverless PostgreSQL, connected via DATABASE_URL.",
    "SQLite (Fallback)",
    "  - Used for local development or as a Vercel serverless fallback (/tmp/egate.db)."
])

# 4. Frontend Technologies
add_slide("Frontend & UI", [
    "Jinja2 (Templating Engine)",
    "  - Renders dynamic HTML using Python variables (e.g. {{ current_user.username }}).",
    "  - Utilizes template inheritance (base.html) for DRY code.",
    "Bootstrap 5",
    "  - CSS framework for responsive design, grid layouts, and cards.",
    "FontAwesome",
    "  - Used for vectorized icons across the dashboard and challenge UI.",
    "Custom CSS & JavaScript",
    "  - Additional styling (cyberpunk-themed elements) and client-side logic."
])

# 5. Security & Authentication
add_slide("Security & Authentication", [
    "Flask-Login",
    "  - Manages user sessions, login state, and cookies.",
    "  - Provides the @login_required decorator to protect routes.",
    "Werkzeug Security",
    "  - Handles secure password hashing (generate_password_hash).",
    "  - Validates login attempts (check_password_hash).",
    "Flask Flash Messages",
    "  - Securely passes one-time alerts (success/error) to the frontend view."
])

# 6. Deployment & Environment
add_slide("Deployment & Config", [
    "Vercel (Serverless Hosting)",
    "  - The platform is designed to be hosted on Vercel's serverless infrastructure.",
    "  - Python functions are executed as serverless lambdas.",
    "Python-dotenv",
    "  - Manages environment variables securely from .env files.",
    "  - Isolates secrets like DATABASE_URL and SECRET_KEY.",
    "Gunicorn / Werkzeug",
    "  - The WSGI HTTP servers used to serve the application."
])

# 7. Summary
add_slide("Summary", [
    "Lightweight & Fast",
    "  - Flask provides exactly what is needed without unnecessary bloat.",
    "Modern Serverless DB",
    "  - Neon PostgreSQL pairs perfectly with Vercel for scalable CTF traffic.",
    "Secure by Design",
    "  - ORM prevents SQLi, Werkzeug hashes passwords, and sessions are signed.",
    "Highly Extensible",
    "  - Blueprint architecture makes adding new CTF challenges easy."
])

prs.save("EGATE_Tech_Stack.pptx")
print("Presentation generated successfully at EGATE_Tech_Stack.pptx")
